#!/usr/bin/env python3
"""
Gestor de Documentos con Interfaz Gráfica - Couce Consulting
Permite seleccionar, subir, eliminar y gestionar documentos manualmente
Con soporte para subcarpetas, panel de log redimensionable, y apertura de archivos
"""

import os
import sys
import time
import json
import shutil
import hashlib
import threading
import subprocess
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
from datetime import datetime
from pathlib import Path

# ============================================
# CONFIGURACIÓN
# ============================================

# Credenciales
GEMINI_KEY = "AIzaSyCBnQe60iOWHMVjJ5L6_uLeGL6gvSmtbX8"
SUPABASE_URL = "https://zejbxsymddswbfabybkh.supabase.co"
SUPABASE_ANON_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InplamJ4c3ltZGRzd2JmYWJ5YmtoIiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzQyMzgyOTAsImV4cCI6MjA4OTgxNDI5MH0.d4CbmZgCKWhqJBw_J_oEYnhrr88vLSf4oGSf48Fm4sg"

# Rutas
RAIZ = "/home/spy/Descargas/couceconsulting/wp-content/uploads/portal-colaboraciones"
CARPETA_ORIGINALES = os.path.join(RAIZ, "originales")
CARPETA_BORRADORES = os.path.join(RAIZ, "borradores")
CARPETA_HISTORICO = os.path.join(RAIZ, "historico")
CARPETA_WORDPRESS = "/home/spy/Descargas/couceconsulting"

# Crear carpetas si no existen
for carpeta in [CARPETA_ORIGINALES, CARPETA_BORRADORES, CARPETA_HISTORICO]:
    os.makedirs(carpeta, exist_ok=True)

# Configuración del sistema
TAMANO_BLOQUE = 3000
LIMITE_MB = 30
EXTENSIONES = (".pdf", ".docx", ".xlsx", ".txt")

# Archivos de registro
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
REGISTRO = os.path.join(BASE_DIR, "pdfs_procesados.json")
LOG_FILE = os.path.join(BASE_DIR, "gestor.log")

# Intentar importar dependencias externas
try:
    import fitz  # pymupdf
    import requests
    import docx as docx_lib
    import openpyxl
except ImportError as e:
    print(f"❌ Error: Falta instalar dependencias: {e}")
    print("Ejecuta: pip install pymupdf requests python-docx openpyxl")
    sys.exit(1)

# ============================================
# FUNCIONES DE UTILIDAD
# ============================================

def calcular_hash_archivo(ruta):
    """Calcula el hash MD5 de un archivo"""
    hash_md5 = hashlib.md5()
    with open(ruta, "rb") as f:
        for chunk in iter(lambda: f.read(4096), b""):
            hash_md5.update(chunk)
    return hash_md5.hexdigest()

def cargar_registro():
    """Carga el registro de archivos procesados."""
    if os.path.exists(REGISTRO):
        try:
            with open(REGISTRO, "r", encoding="utf-8") as f:
                return json.load(f)
        except:
            return {}
    return {}

def guardar_registro(procesados):
    """Guarda el registro de archivos procesados."""
    with open(REGISTRO, "w", encoding="utf-8") as f:
        json.dump(procesados, f, ensure_ascii=False, indent=2)

def sincronizar_con_carpetas(log_func=None):
    """
    Sincroniza el registro con los archivos físicos en ORIGINALES (recursivamente)
    Retorna: (procesados, cantidad_agregados)
    """
    procesados = cargar_registro()
    agregados = 0
    
    def log(mensaje):
        if log_func:
            log_func(mensaje)
        else:
            print(mensaje)
    
    # Verificar si la carpeta existe
    if not os.path.exists(CARPETA_ORIGINALES):
        log(f"⚠️ La carpeta ORIGINALES no existe: {CARPETA_ORIGINALES}")
        return procesados, 0
    
    log(f"📂 Escaneando carpetas recursivamente: {CARPETA_ORIGINALES}")
    
    # Escanear recursivamente todas las subcarpetas
    archivos_encontrados = []
    
    try:
        # Usar os.walk para recorrer recursivamente
        total_encontrados = 0
        for raiz, directorios, archivos in os.walk(CARPETA_ORIGINALES):
            for archivo in archivos:
                if archivo.lower().endswith(EXTENSIONES):
                    ruta_completa = os.path.join(raiz, archivo)
                    archivos_encontrados.append((archivo, ruta_completa))
                    total_encontrados += 1
        
        log(f"📄 Encontrados {total_encontrados} archivos en ORIGINALES y subcarpetas")
        
        # Mostrar estructura de carpetas
        if total_encontrados > 0:
            log(f"📁 Estructura encontrada:")
            carpetas_mostradas = set()
            for nombre, ruta in archivos_encontrados[:10]:
                ruta_rel = os.path.relpath(ruta, CARPETA_ORIGINALES)
                carpeta = os.path.dirname(ruta_rel) if os.path.dirname(ruta_rel) else "(raíz)"
                if carpeta not in carpetas_mostradas:
                    log(f"   📁 {carpeta}/")
                    carpetas_mostradas.add(carpeta)
                log(f"      📄 {nombre}")
        
        for nombre, ruta_completa in archivos_encontrados:
            # Si el archivo no está en el registro, agregarlo
            if nombre not in procesados:
                try:
                    tamano = os.path.getsize(ruta_completa) / (1024 * 1024)
                    hash_archivo = calcular_hash_archivo(ruta_completa)
                    
                    # Obtener la ruta relativa desde CARPETA_ORIGINALES
                    ruta_relativa = os.path.relpath(ruta_completa, CARPETA_ORIGINALES)
                    ruta_borrador = os.path.join(CARPETA_BORRADORES, ruta_relativa)
                    
                    # Crear directorio en BORRADORES si no existe
                    os.makedirs(os.path.dirname(ruta_borrador), exist_ok=True)
                    
                    # Copiar a BORRADORES si no existe
                    if not os.path.exists(ruta_borrador):
                        shutil.copy2(ruta_completa, ruta_borrador)
                        log(f"   💾 Backup creado: {ruta_relativa}")
                    
                    # Obtener fecha del archivo
                    fecha_mod = os.path.getmtime(ruta_completa)
                    
                    procesados[nombre] = {
                        "fecha": fecha_mod,
                        "fecha_str": datetime.fromtimestamp(fecha_mod).strftime('%Y-%m-%d %H:%M:%S'),
                        "hash": hash_archivo,
                        "ruta_original": ruta_completa,
                        "ruta_borrador": ruta_borrador,
                        "ruta_wordpress": obtener_ruta_wordpress(ruta_completa),
                        "tamano_mb": tamano,
                        "ultima_actualizacion": datetime.fromtimestamp(fecha_mod).strftime('%Y-%m-%d %H:%M:%S'),
                        "ruta_relativa": ruta_relativa
                    }
                    agregados += 1
                    log(f"   ✅ Agregado: {nombre} ({tamano:.2f} MB) - {ruta_relativa}")
                    
                except Exception as e:
                    log(f"   ❌ Error procesando {nombre}: {e}")
    
    except Exception as e:
        log(f"❌ Error escaneando carpetas: {e}")
    
    # Guardar si hubo cambios
    if agregados > 0:
        guardar_registro(procesados)
        log(f"✅ Sincronización completada: {agregados} documentos agregados")
    else:
        log(f"✅ No se encontraron documentos nuevos para agregar")
    
    return procesados, agregados

def extraer_texto(ruta):
    """Extrae el texto de un PDF, Word, Excel o TXT."""
    extension = os.path.splitext(ruta)[1].lower()
    
    try:
        if extension == ".pdf":
            doc = fitz.open(ruta)
            texto = ""
            for pagina in doc:
                texto += pagina.get_text() + "\n"
            doc.close()
            return texto
        
        elif extension == ".docx":
            doc = docx_lib.Document(ruta)
            texto = "\n".join([parrafo.text for parrafo in doc.paragraphs if parrafo.text.strip()])
            return texto
        
        elif extension == ".xlsx":
            wb = openpyxl.load_workbook(ruta, data_only=True)
            lineas = []
            for hoja in wb.worksheets:
                lineas.append(f"Hoja: {hoja.title}")
                for fila in hoja.iter_rows(values_only=True):
                    celdas = [str(c) for c in fila if c is not None]
                    if celdas:
                        lineas.append(" | ".join(celdas))
            return "\n".join(lineas)
        
        elif extension == ".txt":
            with open(ruta, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        
        else:
            raise Exception(f"Formato no soportado: {extension}")
            
    except Exception as e:
        raise Exception(f"Error extrayendo texto: {e}")

def trocear_texto(texto):
    """Trocea el texto en bloques para los embeddings."""
    bloques = []
    inicio = 0
    
    while inicio < len(texto):
        fin = inicio + TAMANO_BLOQUE
        
        if fin < len(texto):
            # Intentar cortar por punto o salto de línea
            ultimo_punto = texto.rfind(".", inicio, fin)
            ultimo_salto = texto.rfind("\n", inicio, fin)
            punto_corte = max(ultimo_punto, ultimo_salto)
            
            if punto_corte > inicio + 200:
                fin = punto_corte + 1
        
        bloque = texto[inicio:fin].strip()
        if len(bloque) > 50:
            bloques.append(bloque)
        
        inicio = fin
    
    return bloques

def generar_embedding(texto):
    """Genera el embedding de un texto con Gemini."""
    url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-embedding-001:embedContent?key={GEMINI_KEY}"
    body = {
        "model": "models/gemini-embedding-001",
        "content": {"parts": [{"text": texto[:3000]}]}
    }
    
    try:
        res = requests.post(url, json=body, timeout=30)
        data = res.json()
        
        if "embedding" not in data:
            raise Exception(f"Error en API: {data}")
        
        return data["embedding"]["values"]
    except Exception as e:
        raise Exception(f"Error generando embedding: {e}")

def obtener_ruta_wordpress(ruta_absoluta):
    """Calcula la ruta relativa para WordPress."""
    try:
        ruta_rel = os.path.relpath(ruta_absoluta, CARPETA_WORDPRESS)
        return ruta_rel.replace("\\", "/")
    except ValueError:
        ruta_rel = os.path.relpath(ruta_absoluta, RAIZ)
        return ruta_rel.replace("\\", "/")

def insertar_documento_supabase(contenido, embedding, nombre_pdf, ruta_archivo):
    """Inserta un bloque en Supabase."""
    url = f"{SUPABASE_URL}/rest/v1/documentos"
    headers = {
        "Content-Type": "application/json",
        "apikey": SUPABASE_ANON_KEY,
        "Authorization": f"Bearer {SUPABASE_ANON_KEY}",
        "Prefer": "return=minimal"
    }
    
    body = {
        "contenido": contenido[:2000],
        "nombre_pdf": nombre_pdf,
        "ruta_archivo": ruta_archivo,
        "embedding": f"[{','.join(map(str, embedding))}]"
    }
    
    try:
        res = requests.post(url, json=body, headers=headers, timeout=30)
        if not res.ok:
            raise Exception(f"Error {res.status_code}: {res.text}")
        return True
    except Exception as e:
        raise Exception(f"Error insertando: {e}")

def borrar_documentos_supabase(nombre_pdf):
    """Borra de Supabase todos los bloques de un archivo."""
    url = f"{SUPABASE_URL}/rest/v1/documentos?nombre_pdf=eq.{nombre_pdf}"
    headers = {
        "apikey": SUPABASE_ANON_KEY,
        "Authorization": f"Bearer {SUPABASE_ANON_KEY}"
    }
    
    try:
        res = requests.delete(url, headers=headers, timeout=30)
        if not res.ok:
            raise Exception(f"Error {res.status_code}: {res.text}")
        return True
    except Exception as e:
        raise Exception(f"Error borrando: {e}")

# ============================================
# CLASE DEL GESTOR DE DOCUMENTOS
# ============================================

class GestorDocumentos:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("Couce Consulting - Gestor de Documentos")
        self.root.geometry("1600x950")
        self.root.minsize(1300, 800)
        
        # Variables
        self.procesados = {}
        self.archivos_seleccionados = []
        self.procesando = False
        self.filtro_busqueda = ""
        
        # Configurar colores
        self.colores = {
            'bg_principal': '#1a1a2e',
            'bg_secundario': '#16213e',
            'bg_terciario': '#0f3460',
            'texto': '#e94560',
            'texto_claro': '#ffffff',
            'exito': '#00d8ff',
            'peligro': '#ff6b6b',
            'advertencia': '#ffd93d'
        }
        
        self.setup_ui()
        
        # Sincronizar automáticamente al iniciar
        self.root.after(100, self.sincronizar_inicial)
        
    def sincronizar_inicial(self):
        """Sincronizar con carpetas al iniciar"""
        self.log("=" * 50)
        self.log("🔄 SINCRONIZACIÓN INICIAL")
        self.log(f"📂 Revisando carpeta: {CARPETA_ORIGINALES}")
        
        # Verificar si la carpeta existe
        if os.path.exists(CARPETA_ORIGINALES):
            # Contar archivos recursivamente
            total_archivos = 0
            for raiz, dirs, archivos in os.walk(CARPETA_ORIGINALES):
                for archivo in archivos:
                    if archivo.lower().endswith(EXTENSIONES):
                        total_archivos += 1
            self.log(f"📄 Archivos encontrados en ORIGINALES (incluyendo subcarpetas): {total_archivos}")
        else:
            self.log(f"❌ La carpeta NO EXISTE: {CARPETA_ORIGINALES}")
        
        self.procesados, agregados = sincronizar_con_carpetas(self.log)
        
        if agregados > 0:
            self.log(f"✅ Sincronización completada: {agregados} documentos agregados")
        else:
            self.log("✅ No se encontraron documentos nuevos para sincronizar")
        
        self.log("=" * 50)
        self.cargar_lista_archivos()
        
    def setup_ui(self):
        """Configurar interfaz gráfica"""
        self.root.configure(bg=self.colores['bg_principal'])
        
        # Configurar grid principal con PanedWindow para panel redimensionable
        main_paned = ttk.PanedWindow(self.root, orient=tk.VERTICAL)
        main_paned.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        # Frame superior (contenido principal)
        top_frame = tk.Frame(main_paned, bg=self.colores['bg_principal'])
        main_paned.add(top_frame, weight=3)
        
        # Frame inferior (log redimensionable)
        bottom_frame = tk.Frame(main_paned, bg=self.colores['bg_secundario'])
        main_paned.add(bottom_frame, weight=1)
        
        # Configurar grid dentro de top_frame
        top_frame.grid_rowconfigure(0, weight=0)  # Header
        top_frame.grid_rowconfigure(1, weight=1)  # Contenido principal
        top_frame.grid_columnconfigure(0, weight=0)  # Panel izquierdo
        top_frame.grid_columnconfigure(1, weight=1)  # Panel central
        
        # ========== BARRA SUPERIOR ==========
        header_frame = tk.Frame(top_frame, bg=self.colores['bg_secundario'], height=80)
        header_frame.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 10))
        header_frame.grid_propagate(False)
        
        # Título
        titulo = tk.Label(
            header_frame,
            text="📄 GESTOR DE DOCUMENTOS COUCE CONSULTING",
            font=('Arial', 24, 'bold'),
            bg=self.colores['bg_secundario'],
            fg=self.colores['texto']
        )
        titulo.pack(pady=15)
        
        # ========== PANEL LATERAL IZQUIERDO (CON SCROLL PARA VER TODO) ==========
        panel_izquierdo_container = tk.Frame(top_frame, bg=self.colores['bg_secundario'], width=350)
        panel_izquierdo_container.grid(row=1, column=0, sticky="ns", padx=(0, 10))
        panel_izquierdo_container.grid_propagate(False)
        
        # Canvas y scrollbar para scroll del panel izquierdo
        canvas = tk.Canvas(panel_izquierdo_container, bg=self.colores['bg_secundario'], highlightthickness=0)
        scrollbar = ttk.Scrollbar(panel_izquierdo_container, orient="vertical", command=canvas.yview)
        scrollable_frame = tk.Frame(canvas, bg=self.colores['bg_secundario'])
        
        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        
        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        
        # Configurar el scroll con la rueda del mouse
        def _on_mousewheel(event):
            canvas.yview_scroll(int(-1*(event.delta/120)), "units")
        canvas.bind_all("<MouseWheel>", _on_mousewheel)
        
        panel_izquierdo = scrollable_frame
        
        # Botones de acción
        tk.Label(
            panel_izquierdo,
            text="📁 ACCIONES",
            font=('Arial', 14, 'bold'),
            bg=self.colores['bg_secundario'],
            fg=self.colores['texto']
        ).pack(pady=(20, 10))
        
        # Botón seleccionar archivos
        btn_seleccionar = tk.Button(
            panel_izquierdo,
            text="📂 SELECCIONAR ARCHIVOS",
            command=self.seleccionar_archivos,
            bg=self.colores['texto'],
            fg='white',
            font=('Arial', 11, 'bold'),
            padx=20,
            pady=10,
            cursor='hand2'
        )
        btn_seleccionar.pack(fill=tk.X, padx=20, pady=5)
        
        # Botón seleccionar carpeta
        btn_carpeta = tk.Button(
            panel_izquierdo,
            text="📁 SELECCIONAR CARPETA",
            command=self.seleccionar_carpeta,
            bg='#4ecdc4',
            fg='white',
            font=('Arial', 11, 'bold'),
            padx=20,
            pady=10,
            cursor='hand2'
        )
        btn_carpeta.pack(fill=tk.X, padx=20, pady=5)
        
        # Botón subir archivos
        self.btn_subir = tk.Button(
            panel_izquierdo,
            text="🚀 SUBIR SELECCIONADOS",
            command=self.subir_archivos,
            bg='#2ecc71',
            fg='black',
            font=('Arial', 11, 'bold'),
            padx=20,
            pady=10,
            state=tk.DISABLED,
            cursor='hand2'
        )
        self.btn_subir.pack(fill=tk.X, padx=20, pady=5)
        
        # Separador
        tk.Frame(panel_izquierdo, height=2, bg=self.colores['texto']).pack(fill=tk.X, padx=20, pady=10)
        
        # Botones de gestión
        tk.Label(
            panel_izquierdo,
            text="🔧 GESTIÓN",
            font=('Arial', 14, 'bold'),
            bg=self.colores['bg_secundario'],
            fg=self.colores['texto']
        ).pack(pady=(10, 10))
        
        # Botón ABRIR
        btn_abrir = tk.Button(
            panel_izquierdo,
            text="📂 ABRIR DOCUMENTO",
            command=self.abrir_documento_seleccionado,
            bg='#27ae60',
            fg='white',
            font=('Arial', 11, 'bold'),
            padx=20,
            pady=10,
            cursor='hand2'
        )
        btn_abrir.pack(fill=tk.X, padx=20, pady=5)
        
        # Botón UBICACIÓN
        btn_ubicacion = tk.Button(
            panel_izquierdo,
            text="📍 ABRIR UBICACIÓN",
            command=self.abrir_ubicacion_seleccionado,
            bg='#f39c12',
            fg='white',
            font=('Arial', 11, 'bold'),
            padx=20,
            pady=10,
            cursor='hand2'
        )
        btn_ubicacion.pack(fill=tk.X, padx=20, pady=5)
        
        btn_eliminar = tk.Button(
            panel_izquierdo,
            text="🗑️ ELIMINAR SELECCIONADO",
            command=self.eliminar_seleccionado,
            bg='#e74c3c',
            fg='white',
            font=('Arial', 11, 'bold'),
            padx=20,
            pady=10,
            cursor='hand2'
        )
        btn_eliminar.pack(fill=tk.X, padx=20, pady=5)
        
        btn_sincronizar = tk.Button(
            panel_izquierdo,
            text="🔄 SINCRONIZAR CARPETAS",
            command=self.sincronizar_manual,
            bg='#3498db',
            fg='white',
            font=('Arial', 11, 'bold'),
            padx=20,
            pady=10,
            cursor='hand2'
        )
        btn_sincronizar.pack(fill=tk.X, padx=20, pady=5)
        
        btn_ver_historico = tk.Button(
            panel_izquierdo,
            text="📜 VER HISTÓRICO",
            command=self.ver_historico,
            bg='#f39c12',
            fg='white',
            font=('Arial', 11, 'bold'),
            padx=20,
            pady=10,
            cursor='hand2'
        )
        btn_ver_historico.pack(fill=tk.X, padx=20, pady=5)
        
        btn_refrescar = tk.Button(
            panel_izquierdo,
            text="🔄 REFRESCAR LISTA",
            command=self.cargar_lista_archivos,
            bg='#3498db',
            fg='white',
            font=('Arial', 11, 'bold'),
            padx=20,
            pady=10,
            cursor='hand2'
        )
        btn_refrescar.pack(fill=tk.X, padx=20, pady=5)
        
        btn_ver_logs = tk.Button(
            panel_izquierdo,
            text="📋 VER LOGS",
            command=self.ver_logs,
            bg='#95a5a6',
            fg='white',
            font=('Arial', 11, 'bold'),
            padx=20,
            pady=10,
            cursor='hand2'
        )
        btn_ver_logs.pack(fill=tk.X, padx=20, pady=5)
        
        # ========== BUSCADOR ==========
        tk.Label(
            panel_izquierdo,
            text="🔍 BUSCADOR",
            font=('Arial', 14, 'bold'),
            bg=self.colores['bg_secundario'],
            fg=self.colores['texto']
        ).pack(pady=(20, 10))
        
        frame_buscador = tk.Frame(panel_izquierdo, bg=self.colores['bg_secundario'])
        frame_buscador.pack(fill=tk.X, padx=20, pady=5)
        
        tk.Label(
            frame_buscador,
            text="Buscar:",
            font=('Arial', 10),
            bg=self.colores['bg_secundario'],
            fg=self.colores['texto_claro']
        ).pack(side=tk.LEFT, padx=(0, 5))
        
        self.entry_buscar = tk.Entry(
            frame_buscador,
            font=('Arial', 10),
            bg=self.colores['bg_terciario'],
            fg=self.colores['texto_claro'],
            insertbackground=self.colores['texto_claro']
        )
        self.entry_buscar.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.entry_buscar.bind('<KeyRelease>', self.filtrar_lista)
        
        btn_limpiar = tk.Button(
            frame_buscador,
            text="✖",
            command=self.limpiar_buscador,
            bg=self.colores['peligro'],
            fg='white',
            font=('Arial', 8, 'bold'),
            width=2,
            cursor='hand2'
        )
        btn_limpiar.pack(side=tk.RIGHT, padx=(5, 0))
        
        # ========== ESTADÍSTICAS ==========
        tk.Label(
            panel_izquierdo,
            text="📊 ESTADÍSTICAS",
            font=('Arial', 14, 'bold'),
            bg=self.colores['bg_secundario'],
            fg=self.colores['texto']
        ).pack(pady=(20, 10))
        
        self.stats_frame = tk.Frame(panel_izquierdo, bg=self.colores['bg_secundario'])
        self.stats_frame.pack(fill=tk.X, padx=20, pady=5)
        
        self.lbl_total = tk.Label(
            self.stats_frame,
            text="Total: 0",
            font=('Arial', 11, 'bold'),
            bg=self.colores['bg_secundario'],
            fg=self.colores['exito']
        )
        self.lbl_total.pack(anchor=tk.W, pady=3)
        
        self.lbl_tamano = tk.Label(
            self.stats_frame,
            text="Tamaño: 0 MB",
            font=('Arial', 11, 'bold'),
            bg=self.colores['bg_secundario'],
            fg=self.colores['exito']
        )
        self.lbl_tamano.pack(anchor=tk.W, pady=3)
        
        self.lbl_filtro = tk.Label(
            self.stats_frame,
            text="",
            font=('Arial', 9),
            bg=self.colores['bg_secundario'],
            fg=self.colores['advertencia']
        )
        self.lbl_filtro.pack(anchor=tk.W, pady=2)
        
        # Espacio final para que no quede pegado
        tk.Frame(panel_izquierdo, height=20, bg=self.colores['bg_secundario']).pack()
        
        # ========== PANEL CENTRAL ==========
        panel_central = tk.Frame(top_frame, bg=self.colores['bg_terciario'])
        panel_central.grid(row=1, column=1, sticky="nsew")
        
        # Pestañas
        self.notebook = ttk.Notebook(panel_central)
        self.notebook.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        # Pestaña 1: Archivos en Supabase
        self.tab_subidos = tk.Frame(self.notebook, bg=self.colores['bg_terciario'])
        self.notebook.add(self.tab_subidos, text="✅ Documentos Activos")
        
        # Lista de archivos
        frame_lista = tk.Frame(self.tab_subidos, bg=self.colores['bg_terciario'])
        frame_lista.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        # Treeview con columnas ajustadas
        columns = ('nombre', 'ruta', 'fecha', 'tamano', 'estado')
        self.tree = ttk.Treeview(frame_lista, columns=columns, show='tree headings', height=25)
        
        self.tree.heading('#0', text='')
        self.tree.heading('nombre', text='📄 Nombre del Documento')
        self.tree.heading('ruta', text='📁 Ruta')
        self.tree.heading('fecha', text='📅 Última Actualización')
        self.tree.heading('tamano', text='💾 Tamaño')
        self.tree.heading('estado', text='✓ Estado')
        
        # Configurar columnas
        self.tree.column('#0', width=0, stretch=False)
        self.tree.column('nombre', width=280, minwidth=200, anchor='w')
        self.tree.column('ruta', width=450, minwidth=300, anchor='w')
        self.tree.column('fecha', width=150, minwidth=120, anchor='center')
        self.tree.column('tamano', width=100, minwidth=80, anchor='center')
        self.tree.column('estado', width=120, minwidth=100, anchor='center')
        
        # Scrollbars
        scrollbar_y = ttk.Scrollbar(frame_lista, orient=tk.VERTICAL, command=self.tree.yview)
        scrollbar_x = ttk.Scrollbar(frame_lista, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=scrollbar_y.set, xscrollcommand=scrollbar_x.set)
        
        self.tree.grid(row=0, column=0, sticky="nsew")
        scrollbar_y.grid(row=0, column=1, sticky="ns")
        scrollbar_x.grid(row=1, column=0, sticky="ew")
        
        frame_lista.grid_rowconfigure(0, weight=1)
        frame_lista.grid_columnconfigure(0, weight=1)
        
        # Bind doble clic
        self.tree.bind('<Double-Button-1>', self.on_double_click)
        
        # Pestaña 2: Archivos seleccionados
        self.tab_seleccionados = tk.Frame(self.notebook, bg=self.colores['bg_terciario'])
        self.notebook.add(self.tab_seleccionados, text="📂 Archivos a Subir")
        
        # Lista de archivos seleccionados
        frame_seleccionados = tk.Frame(self.tab_seleccionados, bg=self.colores['bg_terciario'])
        frame_seleccionados.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        self.lista_seleccionados = tk.Listbox(
            frame_seleccionados,
            bg=self.colores['bg_secundario'],
            fg=self.colores['texto_claro'],
            font=('Courier', 10),
            selectmode=tk.EXTENDED
        )
        self.lista_seleccionados.pack(fill=tk.BOTH, expand=True)
        
        # Barra de progreso
        self.progress_bar = ttk.Progressbar(
            panel_central,
            mode='indeterminate',
            length=100
        )
        
        # ========== PANEL INFERIOR (LOG) ==========
        log_titulo_frame = tk.Frame(bottom_frame, bg=self.colores['bg_secundario'])
        log_titulo_frame.pack(fill=tk.X, padx=5, pady=2)
        
        log_titulo = tk.Label(
            log_titulo_frame,
            text="📋 REGISTRO DE ACTIVIDAD (Arrastra el borde superior para redimensionar)",
            font=('Arial', 9, 'bold'),
            bg=self.colores['bg_secundario'],
            fg=self.colores['texto']
        )
        log_titulo.pack(side=tk.LEFT)
        
        # Log de actividad
        self.log_text = scrolledtext.ScrolledText(
            bottom_frame,
            height=8,
            bg=self.colores['bg_terciario'],
            fg=self.colores['texto_claro'],
            font=('Courier', 9),
            wrap=tk.WORD,
            state=tk.DISABLED
        )
        self.log_text.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        self.log("✅ Gestor de documentos iniciado")
        self.log("📁 Estructura de carpetas:")
        self.log(f"   📂 Originales: {CARPETA_ORIGINALES}")
        self.log(f"   📂 Borradores: {CARPETA_BORRADORES}")
        self.log(f"   📂 Histórico: {CARPETA_HISTORICO}")
        
    def log(self, mensaje):
        """Agregar mensaje al log"""
        timestamp = datetime.now().strftime('%H:%M:%S')
        
        self.log_text.config(state=tk.NORMAL)
        self.log_text.insert(tk.END, f"[{timestamp}] {mensaje}\n")
        self.log_text.see(tk.END)
        self.log_text.config(state=tk.DISABLED)
        
        try:
            with open(LOG_FILE, 'a', encoding='utf-8') as f:
                f.write(f"[{timestamp}] {mensaje}\n")
        except:
            pass
    
    def obtener_documento_seleccionado(self):
        """Obtiene el documento seleccionado"""
        seleccion = self.tree.selection()
        if not seleccion:
            return None, None
        item = self.tree.item(seleccion[0])
        nombre = item['values'][0]
        info = self.procesados.get(nombre, {})
        return nombre, info
    
    def abrir_documento_seleccionado(self):
        """Abrir documento seleccionado"""
        nombre, info = self.obtener_documento_seleccionado()
        if not nombre:
            messagebox.showinfo("Info", "Selecciona un documento para abrir")
            return
        
        ruta = info.get('ruta_original', '')
        if not os.path.exists(ruta):
            messagebox.showerror("Error", f"El archivo no existe:\n{ruta}")
            return
        
        try:
            if sys.platform == 'win32':
                os.startfile(ruta)
            elif sys.platform == 'darwin':
                subprocess.run(['open', ruta])
            else:
                subprocess.run(['xdg-open', ruta])
            self.log(f"📂 Abierto: {nombre}")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo abrir el archivo:\n{e}")
    
    def abrir_ubicacion_seleccionado(self):
        """Abrir ubicación del documento"""
        nombre, info = self.obtener_documento_seleccionado()
        if not nombre:
            messagebox.showinfo("Info", "Selecciona un documento para ver su ubicación")
            return
        
        ruta = info.get('ruta_original', '')
        if not os.path.exists(ruta):
            messagebox.showerror("Error", f"El archivo no existe:\n{ruta}")
            return
        
        carpeta = os.path.dirname(ruta)
        try:
            if sys.platform == 'win32':
                os.startfile(carpeta)
            elif sys.platform == 'darwin':
                subprocess.run(['open', carpeta])
            else:
                subprocess.run(['xdg-open', carpeta])
            self.log(f"📍 Abierta ubicación: {carpeta}")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo abrir la ubicación:\n{e}")
    
    def on_double_click(self, event):
        """Maneja doble clic"""
        self.abrir_documento_seleccionado()
    
    def filtrar_lista(self, event=None):
        """Filtrar lista por nombre"""
        self.filtro_busqueda = self.entry_buscar.get().lower().strip()
        self.cargar_lista_archivos()
    
    def limpiar_buscador(self):
        """Limpiar buscador"""
        self.entry_buscar.delete(0, tk.END)
        self.filtro_busqueda = ""
        self.cargar_lista_archivos()
    
    def sincronizar_manual(self):
        """Sincronizar manualmente"""
        self.log("=" * 50)
        self.log("🔄 INICIANDO SINCRONIZACIÓN MANUAL")
        
        self.procesados, agregados = sincronizar_con_carpetas(self.log)
        
        if agregados > 0:
            self.log(f"✅ Sincronización completada: {agregados} documentos agregados")
            messagebox.showinfo("Sincronización", f"Sincronización completada\n\nSe agregaron {agregados} documentos nuevos al registro.")
        else:
            self.log("✅ No se encontraron documentos nuevos")
            messagebox.showinfo("Sincronización", "No se encontraron documentos nuevos para agregar.")
        
        self.log("=" * 50)
        self.cargar_lista_archivos()
        
    def seleccionar_archivos(self):
        """Seleccionar archivos"""
        archivos = filedialog.askopenfilenames(
            title="Seleccionar archivos",
            filetypes=[
                ("Documentos", "*.pdf *.docx *.xlsx *.txt"),
                ("PDF", "*.pdf"),
                ("Word", "*.docx"),
                ("Excel", "*.xlsx"),
                ("Texto", "*.txt"),
                ("Todos", "*.*")
            ]
        )
        
        if archivos:
            self.archivos_seleccionados = list(archivos)
            self.actualizar_lista_seleccionados()
            self.btn_subir.config(state=tk.NORMAL)
            self.log(f"📂 Seleccionados {len(archivos)} archivos")
            
    def seleccionar_carpeta(self):
        """Seleccionar carpeta"""
        carpeta = filedialog.askdirectory(title="Seleccionar carpeta")
        
        if carpeta:
            archivos = []
            for ext in EXTENSIONES:
                archivos.extend(Path(carpeta).rglob(f"*{ext}"))
            
            self.archivos_seleccionados = [str(a) for a in archivos]
            self.actualizar_lista_seleccionados()
            self.btn_subir.config(state=tk.NORMAL)
            self.log(f"📁 Seleccionados {len(archivos)} archivos de {carpeta}")
            
    def actualizar_lista_seleccionados(self):
        """Actualizar lista de seleccionados"""
        self.lista_seleccionados.delete(0, tk.END)
        for archivo in self.archivos_seleccionados:
            nombre = os.path.basename(archivo)
            try:
                tamano = os.path.getsize(archivo) / (1024 * 1024)
                self.lista_seleccionados.insert(tk.END, f"{nombre} ({tamano:.2f} MB)")
            except:
                self.lista_seleccionados.insert(tk.END, f"{nombre} (Error al leer)")
            
    def subir_archivos(self):
        """Subir archivos"""
        if not self.archivos_seleccionados:
            messagebox.showwarning("Advertencia", "No hay archivos seleccionados")
            return
            
        if self.procesando:
            messagebox.showinfo("Info", "Ya hay un proceso en curso")
            return
            
        if not messagebox.askyesno("Confirmar", f"¿Subir {len(self.archivos_seleccionados)} archivos?"):
            return
            
        self.procesando = True
        self.btn_subir.config(state=tk.DISABLED)
        self.progress_bar.pack(fill=tk.X, padx=10, pady=5)
        self.progress_bar.start()
        
        threading.Thread(target=self._procesar_subida, daemon=True).start()
        
    def _procesar_subida(self):
        """Procesar subida"""
        exitos = 0
        fallos = 0
        
        for archivo in self.archivos_seleccionados:
            nombre = os.path.basename(archivo)
            self.root.after(0, self.log, f"📄 Procesando: {nombre}")
            
            try:
                tamano_mb = os.path.getsize(archivo) / (1024 * 1024)
                if tamano_mb > LIMITE_MB:
                    self.root.after(0, self.log, f"   ⚠️ {nombre} - Demasiado grande ({tamano_mb:.1f} MB)")
                    fallos += 1
                    continue
                
                texto = extraer_texto(archivo)
                if not texto.strip():
                    self.root.after(0, self.log, f"   ⚠️ {nombre} - Sin texto extraíble")
                    fallos += 1
                    continue
                
                bloques = trocear_texto(texto)
                self.root.after(0, self.log, f"   📊 {nombre} - {len(bloques)} bloques")
                
                ruta_wp = obtener_ruta_wordpress(archivo)
                todos_exitos = True
                
                for i, bloque in enumerate(bloques):
                    try:
                        embedding = generar_embedding(bloque)
                        insertar_documento_supabase(bloque, embedding, nombre, ruta_wp)
                        if (i + 1) % 10 == 0:
                            self.root.after(0, self.log, f"   ✅ {i+1}/{len(bloques)} bloques")
                    except Exception as e:
                        self.root.after(0, self.log, f"   ❌ Error bloque {i+1}: {e}")
                        todos_exitos = False
                        break
                
                if todos_exitos:
                    ruta_original_existente = None
                    ruta_relativa = None
                    
                    for raiz, dirs, archs in os.walk(CARPETA_ORIGINALES):
                        if nombre in archs:
                            ruta_original_existente = os.path.join(raiz, nombre)
                            ruta_relativa = os.path.relpath(ruta_original_existente, CARPETA_ORIGINALES)
                            break
                    
                    if ruta_relativa:
                        ruta_original = os.path.join(CARPETA_ORIGINALES, ruta_relativa)
                        ruta_borrador = os.path.join(CARPETA_BORRADORES, ruta_relativa)
                    else:
                        ruta_original = os.path.join(CARPETA_ORIGINALES, nombre)
                        ruta_borrador = os.path.join(CARPETA_BORRADORES, nombre)
                        ruta_relativa = nombre
                    
                    os.makedirs(os.path.dirname(ruta_original), exist_ok=True)
                    os.makedirs(os.path.dirname(ruta_borrador), exist_ok=True)
                    
                    es_actualizacion = os.path.exists(ruta_original)
                    
                    if es_actualizacion:
                        self.root.after(0, self.log, f"   🔄 Actualizando documento existente...")
                        
                        nombre_sin_ext, ext = os.path.splitext(nombre)
                        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                        nombre_historico = f"{nombre_sin_ext}_v{timestamp}{ext}"
                        ruta_historico = os.path.join(CARPETA_HISTORICO, nombre_historico)
                        shutil.move(ruta_original, ruta_historico)
                        self.root.after(0, self.log, f"   📜 Versión anterior archivada")
                        
                        shutil.move(archivo, ruta_original)
                        self.root.after(0, self.log, f"   📁 Nueva versión guardada en ORIGINALES")
                        
                        shutil.copy2(ruta_original, ruta_borrador)
                        self.root.after(0, self.log, f"   💾 Backup actualizado")
                        
                    else:
                        self.root.after(0, self.log, f"   📄 Nuevo documento...")
                        
                        shutil.move(archivo, ruta_original)
                        self.root.after(0, self.log, f"   📁 Guardado en ORIGINALES")
                        
                        shutil.copy2(ruta_original, ruta_borrador)
                        self.root.after(0, self.log, f"   💾 Backup creado")
                    
                    self.procesados[nombre] = {
                        "fecha": datetime.now().timestamp(),
                        "fecha_str": datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                        "hash": calcular_hash_archivo(ruta_original),
                        "ruta_original": ruta_original,
                        "ruta_borrador": ruta_borrador,
                        "ruta_wordpress": ruta_wp,
                        "tamano_mb": tamano_mb,
                        "ultima_actualizacion": datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
                        "ruta_relativa": ruta_relativa
                    }
                    guardar_registro(self.procesados)
                    exitos += 1
                    self.root.after(0, self.log, f"   ✅ {nombre} completado")
                else:
                    fallos += 1
                    
            except Exception as e:
                self.root.after(0, self.log, f"   ❌ Error: {e}")
                fallos += 1
        
        self.root.after(0, self.log, f"📊 Subida: {exitos} éxitos, {fallos} fallos")
        self.root.after(0, self._finalizar_subida)
        
    def _finalizar_subida(self):
        """Finalizar subida"""
        self.procesando = False
        self.progress_bar.stop()
        self.progress_bar.pack_forget()
        self.btn_subir.config(state=tk.NORMAL)
        self.cargar_lista_archivos()
        self.archivos_seleccionados = []
        self.actualizar_lista_seleccionados()
        
    def cargar_lista_archivos(self):
        """Cargar lista de archivos con filtro"""
        for item in self.tree.get_children():
            self.tree.delete(item)
        
        total = 0
        tamano_total = 0.0
        total_mostrados = 0
        
        for nombre, info in self.procesados.items():
            if self.filtro_busqueda and self.filtro_busqueda not in nombre.lower():
                continue
            
            ruta = info.get('ruta_original', '')
            if not os.path.exists(ruta):
                continue
            
            try:
                tamano_bytes = os.path.getsize(ruta)
                tamano_mb = tamano_bytes / (1024 * 1024)
                tamano_total += tamano_mb
            except:
                tamano_mb = info.get('tamano_mb', 0)
                tamano_total += tamano_mb
            
            fecha = info.get('fecha_str', '')
            if not fecha:
                try:
                    fecha_mod = os.path.getmtime(ruta)
                    fecha = datetime.fromtimestamp(fecha_mod).strftime('%Y-%m-%d %H:%M:%S')
                except:
                    fecha = "Fecha desconocida"
            
            ruta_relativa = info.get('ruta_relativa', '')
            if not ruta_relativa:
                try:
                    ruta_relativa = os.path.relpath(ruta, CARPETA_ORIGINALES)
                except:
                    ruta_relativa = nombre
            
            estado = "✅ Activo"
            if os.path.exists(info.get('ruta_borrador', '')):
                estado = "✅ Activo + Backup"
            
            self.tree.insert('', tk.END, values=(
                nombre,
                ruta_relativa,
                fecha,
                f"{tamano_mb:.2f} MB",
                estado
            ))
            total_mostrados += 1
            total += 1
        
        total_documentos = len(self.procesados)
        self.lbl_total.config(text=f"📄 Total: {total_documentos} documentos")
        self.lbl_tamano.config(text=f"💾 Tamaño: {tamano_total:.2f} MB")
        
        if self.filtro_busqueda:
            self.lbl_filtro.config(text=f"🔍 Mostrando {total_mostrados} de {total_documentos} (filtro: '{self.filtro_busqueda}')")
        else:
            self.lbl_filtro.config(text="")
        
        self.log(f"📊 Lista: {total_mostrados} documentos mostrados, {tamano_total:.2f} MB total")
        
    def eliminar_seleccionado(self):
        """Eliminar archivo seleccionado"""
        seleccion = self.tree.selection()
        if not seleccion:
            messagebox.showinfo("Info", "Selecciona un archivo")
            return
            
        if not messagebox.askyesno("Confirmar", "¿Eliminar este archivo completamente?"):
            return
            
        item = self.tree.item(seleccion[0])
        nombre = item['values'][0]
        info = self.procesados.get(nombre, {})
        
        try:
            if borrar_documentos_supabase(nombre):
                ruta_original = info.get('ruta_original', '')
                ruta_borrador = info.get('ruta_borrador', '')
                
                historico_eliminados = 0
                nombre_sin_ext = os.path.splitext(nombre)[0]
                if os.path.exists(CARPETA_HISTORICO):
                    for archivo_historico in os.listdir(CARPETA_HISTORICO):
                        if archivo_historico.startswith(nombre_sin_ext) and archivo_historico.lower().endswith(EXTENSIONES):
                            try:
                                os.remove(os.path.join(CARPETA_HISTORICO, archivo_historico))
                                historico_eliminados += 1
                            except:
                                pass
                
                if os.path.exists(ruta_original):
                    os.remove(ruta_original)
                    self.log(f"   🗑️ Eliminado original")
                
                if os.path.exists(ruta_borrador):
                    os.remove(ruta_borrador)
                    self.log(f"   🗑️ Eliminado backup")
                
                if historico_eliminados > 0:
                    self.log(f"   🗑️ Eliminadas {historico_eliminados} versiones históricas")
                
                if nombre in self.procesados:
                    del self.procesados[nombre]
                    guardar_registro(self.procesados)
                
                self.log(f"🗑️ Eliminado: {nombre}")
                self.cargar_lista_archivos()
                messagebox.showinfo("Éxito", f"{nombre} eliminado")
            else:
                messagebox.showerror("Error", f"No se pudo eliminar {nombre} de Supabase")
        except Exception as e:
            self.log(f"❌ Error: {e}")
            messagebox.showerror("Error", str(e))
    
    def ver_historico(self):
        """Ver histórico"""
        ventana_historico = tk.Toplevel(self.root)
        ventana_historico.title("Histórico de Documentos")
        ventana_historico.geometry("900x600")
        
        frame = tk.Frame(ventana_historico, bg=self.colores['bg_principal'])
        frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        titulo = tk.Label(
            frame,
            text="📜 HISTÓRICO DE DOCUMENTOS",
            font=('Arial', 16, 'bold'),
            bg=self.colores['bg_principal'],
            fg=self.colores['texto']
        )
        titulo.pack(pady=10)
        
        frame_lista = tk.Frame(frame, bg=self.colores['bg_secundario'])
        frame_lista.pack(fill=tk.BOTH, expand=True)
        
        columns = ('nombre', 'fecha', 'tamano')
        tree_historico = ttk.Treeview(frame_lista, columns=columns, show='tree headings', height=20)
        
        tree_historico.heading('#0', text='')
        tree_historico.heading('nombre', text='Nombre del Archivo')
        tree_historico.heading('fecha', text='Fecha')
        tree_historico.heading('tamano', text='Tamaño')
        
        tree_historico.column('#0', width=0)
        tree_historico.column('nombre', width=500)
        tree_historico.column('fecha', width=200, anchor='center')
        tree_historico.column('tamano', width=100, anchor='center')
        
        scrollbar = ttk.Scrollbar(frame_lista, orient=tk.VERTICAL, command=tree_historico.yview)
        tree_historico.configure(yscrollcommand=scrollbar.set)
        
        tree_historico.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        contador = 0
        if os.path.exists(CARPETA_HISTORICO):
            for archivo in sorted(os.listdir(CARPETA_HISTORICO), reverse=True):
                if archivo.lower().endswith(EXTENSIONES):
                    ruta = os.path.join(CARPETA_HISTORICO, archivo)
                    try:
                        tamano = os.path.getsize(ruta) / (1024 * 1024)
                        fecha = datetime.fromtimestamp(os.path.getmtime(ruta)).strftime('%Y-%m-%d %H:%M:%S')
                        tree_historico.insert('', tk.END, values=(archivo, fecha, f"{tamano:.2f} MB"))
                        contador += 1
                    except:
                        pass
        
        lbl_total = tk.Label(
            frame,
            text=f"Total: {contador} versiones históricas",
            font=('Arial', 10),
            bg=self.colores['bg_principal'],
            fg=self.colores['texto_claro']
        )
        lbl_total.pack(pady=10)
        
        btn_cerrar = tk.Button(
            frame,
            text="CERRAR",
            command=ventana_historico.destroy,
            bg=self.colores['texto'],
            fg='white',
            font=('Arial', 10, 'bold'),
            padx=20,
            pady=5
        )
        btn_cerrar.pack(pady=10)
        
    def ver_logs(self):
        """Ver logs"""
        ventana_logs = tk.Toplevel(self.root)
        ventana_logs.title("Logs del Sistema")
        ventana_logs.geometry("800x500")
        
        texto_logs = scrolledtext.ScrolledText(
            ventana_logs,
            wrap=tk.WORD,
            font=('Courier', 10),
            state=tk.DISABLED
        )
        texto_logs.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        try:
            with open(LOG_FILE, 'r') as f:
                texto_logs.config(state=tk.NORMAL)
                texto_logs.insert(tk.END, f.read())
                texto_logs.config(state=tk.DISABLED)
        except:
            texto_logs.config(state=tk.NORMAL)
            texto_logs.insert(tk.END, "No hay logs disponibles")
            texto_logs.config(state=tk.DISABLED)
            
    def run(self):
        """Ejecutar aplicación"""
        self.root.mainloop()


def main():
    """Función principal"""
    try:
        app = GestorDocumentos()
        app.run()
    except Exception as e:
        print(f"Error: {e}")
        import traceback
        traceback.print_exc()
        input("Presiona Enter para salir...")


if __name__ == "__main__":
    main()